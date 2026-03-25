"""
BOM exporters – Excel (xlsx) and CSV output.
"""

import csv
import logging
from pathlib import Path
from typing import Union

from .generator import BOM

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------

def export_csv(bom: BOM, path: Union[str, Path]):
    """Export BOM line items to a CSV file."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    with open(path, "w", newline="", encoding="utf-8") as f:
        if not bom.line_items:
            f.write("No items in BOM.\n")
            return

        writer = csv.DictWriter(f, fieldnames=bom.line_items[0].as_dict().keys())
        writer.writeheader()
        for item in bom.line_items:
            writer.writerow(item.as_dict())

        f.write("\n")
        f.write("MATERIAL SUMMARY\n")
        mat_fields = bom.material_summary[0].as_dict().keys() if bom.material_summary else []
        if mat_fields:
            writer2 = csv.DictWriter(f, fieldnames=mat_fields)
            writer2.writeheader()
            for ms in bom.material_summary:
                writer2.writerow(ms.as_dict())

    logger.info("BOM exported to CSV: %s", path)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_excel(bom: BOM, path: Union[str, Path]):
    """Export BOM to an Excel workbook with formatting."""
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side, GradientFill,
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed – falling back to CSV export.")
        csv_path = Path(path).with_suffix(".csv")
        export_csv(bom, csv_path)
        return

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # ----- Styles -----
    hdr_font   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    hdr_fill   = PatternFill("solid", fgColor="1F4E79")
    sub_fill   = PatternFill("solid", fgColor="2E75B6")
    alt_fill   = PatternFill("solid", fgColor="D6E4F0")
    center     = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left       = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin       = Side(style="thin", color="AAAAAA")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_font = Font(name="Calibri", bold=True, size=14, color="1F4E79")

    # ----- BOM Sheet -----
    ws = wb.active
    ws.title = "PP-BOM"

    # Title row
    ws.merge_cells("A1:J1")
    ws["A1"] = f"PARTS & PIECES BILL OF MATERIALS – {bom.project_name}"
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    # Project info row
    ws.merge_cells("A2:E2")
    ws["A2"] = f"Project No: {bom.project_number}    Drawn by: {bom.drawn_by}    Date: {bom.date}"
    ws["A2"].font = Font(name="Calibri", italic=True, size=10)
    ws.row_dimensions[2].height = 20

    # Header row
    headers = list(bom.line_items[0].as_dict().keys()) if bom.line_items else [
        "Item No", "Tag", "Description", "Size", "Material",
        "Gauge (mm)", "Quantity", "Unit", "Area (m²)", "Notes",
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[4].height = 25

    # Data rows
    col_widths = [6, 12, 30, 22, 22, 10, 10, 6, 10, 20]
    for row_idx, item in enumerate(bom.line_items, start=5):
        for col_idx, (key, val) in enumerate(item.as_dict().items(), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            cell.alignment = center if col_idx not in (3, 4, 5) else left
            if row_idx % 2 == 0:
                cell.fill = alt_fill
        ws.row_dimensions[row_idx].height = 18

    # Column widths
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze pane
    ws.freeze_panes = "A5"

    # Auto-filter
    if bom.line_items:
        last_col = get_column_letter(len(headers))
        last_row = 4 + len(bom.line_items)
        ws.auto_filter.ref = f"A4:{last_col}{last_row}"

    # ----- Material Summary Sheet -----
    ws2 = wb.create_sheet("Material Summary")
    ws2.merge_cells("A1:C1")
    ws2["A1"] = "MATERIAL SUMMARY"
    ws2["A1"].font = title_font
    ws2["A1"].alignment = center
    ws2.row_dimensions[1].height = 28

    mat_headers = ["Material", "Gauge (mm)", "Total Area (m²)"]
    for col_idx, h in enumerate(mat_headers, start=1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = hdr_font
        cell.fill = sub_fill
        cell.alignment = center
        cell.border = border
    ws2.row_dimensions[3].height = 22

    for row_idx, ms in enumerate(bom.material_summary, start=4):
        for col_idx, (key, val) in enumerate(ms.as_dict().items(), start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            cell.alignment = center
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Grand total row
    total_row = 4 + len(bom.material_summary)
    ws2.cell(row=total_row, column=1, value="GRAND TOTAL").font = Font(bold=True)
    grand_total = ws2.cell(row=total_row, column=3, value=round(bom.total_area_sqm(), 3))
    grand_total.font = Font(bold=True)
    for col in range(1, 4):
        ws2.cell(row=total_row, column=col).border = border

    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 18

    wb.save(path)
    logger.info("BOM exported to Excel: %s", path)
