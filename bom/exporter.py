"""
BOM export to Excel (.xlsx) and CSV.
"""

from __future__ import annotations
import csv
import logging
from pathlib import Path
from typing import Union

from .calculator import FullBOM

logger = logging.getLogger(__name__)


def export_csv(bom: FullBOM, path: Union[str, Path]):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        if not bom.items:
            f.write("No items\n")
            return
        w = csv.DictWriter(f, fieldnames=bom.items[0].as_dict().keys())
        w.writeheader()
        for item in bom.items:
            w.writerow(item.as_dict())
        f.write("\n")
        f.write(f"TOTAL WEIGHT (kg),{bom.total_weight_kg:.3f}\n")
        f.write(f"TOTAL COST (USD),{bom.total_cost_usd:.2f}\n")
    logger.info("BOM CSV: %s", path)


def export_excel(bom: FullBOM, path: Union[str, Path]):
    try:
        import openpyxl
        from openpyxl.styles import (Font, PatternFill, Alignment,
                                      Border, Side, GradientFill)
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed – exporting CSV instead.")
        export_csv(bom, Path(path).with_suffix(".csv"))
        return

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # ── Styles ────────────────────────────────────────────────────────────────
    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    alt_fill  = PatternFill("solid", fgColor="D6E4F0")
    sum_fill  = PatternFill("solid", fgColor="E2EFDA")
    hdr_font  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    ttl_font  = Font(name="Calibri", bold=True, size=14, color="1F4E79")
    bold      = Font(name="Calibri", bold=True, size=10)
    norm      = Font(name="Calibri", size=10)
    center    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left      = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin      = Side(style="thin", color="BBBBBB")
    bdr       = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── BOM sheet ─────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Bill of Materials"

    # Title
    ws.merge_cells("A1:K1")
    ws["A1"] = f"BILL OF MATERIALS – {bom.project_name}"
    ws["A1"].font = ttl_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 32

    # Project info
    ws.merge_cells("A2:K2")
    ws["A2"] = (f"Project No: {bom.project_no}    "
                f"Drawn By: {bom.drawn_by}    Date: {bom.date}")
    ws["A2"].font = Font(name="Calibri", italic=True, size=10)
    ws.row_dimensions[2].height = 18

    if not bom.items:
        ws["A4"] = "No items in BOM."
        wb.save(path)
        return

    # Column headers
    headers = list(bom.items[0].as_dict().keys())
    col_widths = [5, 8, 28, 30, 8, 6, 12, 12, 12, 12, 20]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = bdr
    ws.row_dimensions[4].height = 22

    # Data rows
    for ri, item in enumerate(bom.items, start=5):
        for ci, (k, v) in enumerate(item.as_dict().items(), 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = norm
            cell.border = bdr
            cell.alignment = center if ci not in (3, 4, 11) else left
            if ri % 2 == 0:
                cell.fill = alt_fill
        ws.row_dimensions[ri].height = 16

    # Totals row
    tr = 5 + len(bom.items)
    ws.cell(row=tr, column=1, value="TOTALS").font = bold
    for ci in range(1, len(headers)+1):
        ws.cell(row=tr, column=ci).fill = sum_fill
        ws.cell(row=tr, column=ci).border = bdr
    ws.cell(row=tr, column=8, value=round(bom.total_weight_kg, 3)).font = bold
    ws.cell(row=tr, column=10, value=round(bom.total_cost_usd, 2)).font = bold

    # Column widths
    for ci, cw in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = cw

    ws.freeze_panes = "A5"
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A4:{last_col}{tr-1}"

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2.merge_cells("A1:C1")
    ws2["A1"] = "BOM SUMMARY"
    ws2["A1"].font = ttl_font
    ws2["A1"].alignment = center
    ws2.row_dimensions[1].height = 28

    summary_headers = ["Category", "Total Cost ($)", "% of Total"]
    for ci, h in enumerate(summary_headers, 1):
        cell = ws2.cell(row=3, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = bdr

    total = bom.total_cost_usd or 1
    for ri, (cat, cost) in enumerate(bom.summary_by_category().items(), start=4):
        ws2.cell(row=ri, column=1, value=cat).border = bdr
        ws2.cell(row=ri, column=2, value=round(cost, 2)).border = bdr
        ws2.cell(row=ri, column=3, value=f"{cost/total*100:.1f}%").border = bdr
        if ri % 2 == 0:
            for ci in range(1, 4):
                ws2.cell(row=ri, column=ci).fill = alt_fill

    sr = 4 + len(bom.summary_by_category())
    ws2.cell(row=sr, column=1, value="GRAND TOTAL").font = bold
    ws2.cell(row=sr, column=2, value=round(bom.total_cost_usd, 2)).font = bold
    ws2.cell(row=sr, column=3, value="100%").font = bold
    for ci in range(1, 4):
        ws2.cell(row=sr, column=ci).fill = sum_fill
        ws2.cell(row=sr, column=ci).border = bdr

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 12

    ws2.cell(row=sr+2, column=1,
             value=f"Total Weight: {bom.total_weight_kg:.2f} kg").font = bold

    wb.save(path)
    logger.info("BOM Excel: %s", path)
